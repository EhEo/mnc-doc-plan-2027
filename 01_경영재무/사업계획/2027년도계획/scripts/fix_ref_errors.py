# 2027년도 사업계획 엑셀 #REF! 오류 일괄 수정 스크립트
import sys, io, shutil, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.chartsheet import Chartsheet

SRC = '../working/20260615-2027Y BIZ Plan_Rev1.xlsx'
DST = '../working/20260615-2027Y BIZ Plan_Rev1_Fixed.xlsx'

shutil.copy(SRC, DST)
wb = load_workbook(DST, data_only=False)

fixed_count = 0
log = []

def fix(ws, cell_or_coord, new_value, reason=''):
    global fixed_count
    if isinstance(cell_or_coord, str):
        cell = ws[cell_or_coord]
    else:
        cell = cell_or_coord
    old = cell.value
    cell.value = new_value
    fixed_count += 1
    log.append(f"[{ws.title}] {cell.coordinate}: {repr(old)[:60]} -> {repr(new_value)[:40]}  ({reason})")

# ────────────────────────────────────────────────────────────
# 1. 4. L PLAN  (367셀)
# ────────────────────────────────────────────────────────────
ws = wb['4. L PLAN']
print('4. L PLAN 수정 중...')

# 패턴 1: Row 4 (S:BE, 39셀) — 월 헤더를 Row 28 값으로 하드코딩
row28_map = {col: ws.cell(row=28, column=col).value for col in range(19, 58)}
for col in range(19, 58):
    cell = ws.cell(row=4, column=col)
    if cell.value and '#REF!' in str(cell.value):
        fix(ws, cell, row28_map.get(col, 'N/A'), 'Row4 헤더 → Row28 값')

# 패턴 2: Row 30 (36셀) — (prevROW36-#REF!)*$E30 → #REF! = 0
for col in range(19, 57):
    if col in (31, 44):  # AE, AR 합계열 제외
        continue
    cell = ws.cell(row=30, column=col)
    if cell.value and '#REF!' in str(cell.value):
        new_f = str(cell.value).replace('#REF!', '0')
        fix(ws, cell, new_f, 'Row30 감소인원 #REF!→0')

# 패턴 3: Rows 70-74 (180셀) — ($C$N+$D$N)*#REF! → *{col}{122-126}
ROW_TO_REF = {70: 122, 71: 123, 72: 124, 73: 125, 74: 126}
for prow, ref_row in ROW_TO_REF.items():
    for col in range(19, 57):
        if col in (31, 44):
            continue
        cell = ws.cell(row=prow, column=col)
        if cell.value and '#REF!' in str(cell.value):
            col_letter = get_column_letter(col)
            new_f = str(cell.value).replace('#REF!', f'{col_letter}{ref_row}')
            fix(ws, cell, new_f, f'Row{prow} 급여 → *{col_letter}{ref_row}')

# 패턴 4: Row 96 (36셀) — SUM(#REF!)* → SUM({col}36)*
for col in range(19, 57):
    if col in (31, 44):
        continue
    cell = ws.cell(row=96, column=col)
    if cell.value and '#REF!' in str(cell.value):
        col_letter = get_column_letter(col)
        new_f = str(cell.value).replace('#REF!', f'{col_letter}36')
        fix(ws, cell, new_f, f'Row96 식대 → {col_letter}36 (기말인원)')

# 패턴 5: Row 141 (36셀) — S Plan 2025 매출 → 0
for col in range(19, 57):
    if col in (31, 44):
        continue
    cell = ws.cell(row=141, column=col)
    if cell.value and '#REF!' in str(cell.value):
        fix(ws, cell, 0, 'Row141 참조용 매출액 → 0 (2025 삭제)')

# 패턴 6: Row 142 (4셀: R,AE,AR,BE) — X141/#REF! → X141/X36
for col_letter in ('R', 'AE', 'AR', 'BE'):
    col = column_index_from_string(col_letter)
    cell = ws.cell(row=142, column=col)
    if cell.value and '#REF!' in str(cell.value):
        new_f = str(cell.value).replace('#REF!', f'{col_letter}36')
        fix(ws, cell, new_f, f'Row142 인당매출 → /{col_letter}36')

# 패턴 7: Row 143 (36셀) — P PLAN 2025 생산수량 → 0
for col in range(19, 57):
    if col in (31, 44):
        continue
    cell = ws.cell(row=143, column=col)
    if cell.value and '#REF!' in str(cell.value):
        fix(ws, cell, 0, 'Row143 참조용 생산수량 → 0 (2025 삭제)')

print(f'  → 4. L PLAN: {fixed_count}셀 수정')

# ────────────────────────────────────────────────────────────
# 2. 4. L PLAN a  (28셀)
# ────────────────────────────────────────────────────────────
ws = wb['4. L PLAN a']
cnt_before = fixed_count
print('4. L PLAN a 수정 중...')

# C29:C33 — SUM(#REF!)/'1-1.Assumption'!$F$29 → #REF! = 0
for row in range(29, 34):
    cell = ws.cell(row=row, column=3)
    if cell.value and '#REF!' in str(cell.value):
        new_f = str(cell.value).replace('#REF!', '0')
        fix(ws, cell, new_f, 'C29-33 평균급여 #REF!→0')

# G40:Q40 — =#REF! → =F40 (1월 값 전파)
for col in range(7, 18):  # G=7 to Q=17
    cell = ws.cell(row=40, column=col)
    if cell.value and '#REF!' in str(cell.value):
        fix(ws, cell, '=F40', 'G40-Q40 생산인원 → =F40')

# F47:Q47 — '2. S Plan'!#REF!/F46 → 0
for col in range(6, 18):
    cell = ws.cell(row=47, column=col)
    if cell.value and '#REF!' in str(cell.value):
        fix(ws, cell, 0, 'F47-Q47 매출/인원 비율 → 0')

print(f'  → 4. L PLAN a: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 3. 4-1. 인원계획  (51셀)
# ────────────────────────────────────────────────────────────
ws = wb['4-1. 인원계획']
cnt_before = fixed_count
print('4-1. 인원계획 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            fix(ws, cell, 0, '인원계획 #REF! → 0')

print(f'  → 4-1. 인원계획: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 4. 3. P PLAN-2  (71셀)
# ────────────────────────────────────────────────────────────
ws = wb['3. P PLAN-2']
cnt_before = fixed_count
print('3. P PLAN-2 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            v = str(cell.value)
            if v.startswith('='):
                # IFERROR 수식 내 #REF! → 0 (0 제공 시 IFERROR가 0 반환)
                new_f = v.replace('#REF!', '0')
                fix(ws, cell, new_f, 'P PLAN-2 IFERROR #REF!→0')
            else:
                fix(ws, cell, 0, 'P PLAN-2 #REF! → 0')

print(f'  → 3. P PLAN-2: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 5. 3. P PLAN  (16셀)
# ────────────────────────────────────────────────────────────
ws = wb['3. P PLAN']
cnt_before = fixed_count
print('3. P PLAN 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            v = str(cell.value)
            if v.startswith('=IFERROR'):
                new_f = v.replace('#REF!', '0')
                fix(ws, cell, new_f, 'P PLAN IFERROR #REF!→0')
            else:
                fix(ws, cell, 0, 'P PLAN #REF! → 0')

print(f'  → 3. P PLAN: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 6. 10. IS  (16셀)
# ────────────────────────────────────────────────────────────
ws = wb['10. IS']
cnt_before = fixed_count
print('10. IS 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            # (#REF!-P4)/#REF! 구조 → 이전계획 데이터 삭제됨, 0으로 초기화
            fix(ws, cell, 0, 'IS Q열 이전계획 비교 → 0 (이전계획 삭제)')

print(f'  → 10. IS: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 7. AP Plan  (13셀)
# ────────────────────────────────────────────────────────────
ws = wb['AP Plan']
cnt_before = fixed_count
print('AP Plan 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            fix(ws, cell, 0, 'AP Plan M Plan 삭제행 → 0')

print(f'  → AP Plan: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 8. Utilty_Data  (13셀)
# ────────────────────────────────────────────────────────────
ws = wb['Utilty_Data']
cnt_before = fixed_count
print('Utilty_Data 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            fix(ws, cell, 0, 'Utilty_Data S Plan 삭제열 → 0')

print(f'  → Utilty_Data: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 9. 6. O Plan  (5셀)
# ────────────────────────────────────────────────────────────
ws = wb['6. O Plan']
cnt_before = fixed_count
print('6. O Plan 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            fix(ws, cell, 0, 'O Plan S Plan 삭제행 → 0')

print(f'  → 6. O Plan: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 10. 1. Assumption  (2셀) — 정밀 수식 수정
# ────────────────────────────────────────────────────────────
ws = wb['1. Assumption']
cnt_before = fixed_count
print('1. Assumption 수정 중...')

# E12: +'2. S Plan'!#REF! 항목만 제거 (나머지 합산 유지)
cell_e12 = ws['E12']
if cell_e12.value and '#REF!' in str(cell_e12.value):
    new_f = re.sub(r"\+'2\. S Plan'!#REF!", '', str(cell_e12.value))
    new_f = re.sub(r"'2\. S Plan'!#REF!\+", '', new_f)
    new_f = new_f.replace('#REF!', '0')
    fix(ws, cell_e12, new_f, 'E12 Mando 삭제행 제거')

# E14: 복수의 +'2. S Plan'!#REF! 항목 제거
cell_e14 = ws['E14']
if cell_e14.value and '#REF!' in str(cell_e14.value):
    new_f = re.sub(r"\+'2\. S Plan'!#REF!", '', str(cell_e14.value))
    new_f = re.sub(r"'2\. S Plan'!#REF!\+", '', new_f)
    new_f = new_f.replace('#REF!', '0')
    fix(ws, cell_e14, new_f, 'E14 A사 삭제행 제거')

print(f'  → 1. Assumption: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 11. 1. Master Plan  (2셀)
# ────────────────────────────────────────────────────────────
ws = wb['1. Master Plan']
cnt_before = fixed_count
print('1. Master Plan 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        if cell.value and '#REF!' in str(cell.value):
            fix(ws, cell, 0, 'Master Plan I Plan 삭제행 → 0')

print(f'  → 1. Master Plan: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 12. B.BasicInfomation  (1셀)
# ────────────────────────────────────────────────────────────
ws = wb['B.BasicInfomation']
cnt_before = fixed_count
print('B.BasicInfomation 수정 중...')

for row_data in ws.iter_rows():
    for cell in row_data:
        v = cell.value
        if v == '#REF!' or (v and isinstance(v, str) and '#REF!' in v):
            fix(ws, cell, 0, 'B.BasicInfomation #REF! → 0')

print(f'  → B.BasicInfomation: {fixed_count - cnt_before}셀 수정')

# ────────────────────────────────────────────────────────────
# 저장 및 검증
# ────────────────────────────────────────────────────────────
wb.save(DST)
print(f'\n=== 저장 완료: {DST} ===')
print(f'총 {fixed_count}개 셀 수정')

# 잔여 #REF! 확인
print('\n=== 잔여 #REF! 검증 ===')
wb2 = load_workbook(DST, data_only=False)
remaining = 0
for sn in wb2.sheetnames:
    ws_v = wb2[sn]
    if isinstance(ws_v, Chartsheet):
        continue
    for row_data in ws_v.iter_rows():
        for cell in row_data:
            if cell.value and '#REF!' in str(cell.value):
                remaining += 1
                print(f'  잔여: [{sn}] {cell.coordinate}: {repr(cell.value)[:60]}')

if remaining == 0:
    print('  모든 #REF! 오류가 제거되었습니다.')
else:
    print(f'  잔여 {remaining}개 — 추가 확인 필요')

# 수정 로그 저장
with open('output/_ref_fix_log.txt', 'w', encoding='utf-8') as f:
    f.write(f'총 {fixed_count}개 수정\n\n')
    for line in log:
        f.write(line + '\n')
print(f'\n수정 로그: output/_ref_fix_log.txt')
