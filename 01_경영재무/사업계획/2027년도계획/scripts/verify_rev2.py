# Rev2 보고서 수치 검증 스크립트
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from openpyxl import load_workbook

wb = load_workbook('../working/20260615-2027Y BIZ Plan_Rev2.xlsx', data_only=True)
ws = wb['10. IS']

sales = ws['N4'].value
cogs  = ws['N5'].value
gp    = ws['N8'].value
sga   = ws['N9'].value
op    = ws['N11'].value
ni    = ws['N18'].value

print('=== IS N열 검증 ===')
print('매출:     ${:,.0f}'.format(sales))
print('COGS:    ${:,.0f} ({:.2f}%)'.format(cogs, cogs/sales*100))
print('GP:      ${:,.0f} ({:.2f}%)'.format(gp, gp/sales*100))
print('SGA:     ${:,.0f}'.format(sga))
print('영업이익: ${:,.0f} ({:.2f}%)'.format(op, op/sales*100))
print('순이익:  ${:,.0f}'.format(ni))

feb_s = ws['C4'].value
feb_c = ws['C5'].value
print('2월 COGS%: {:.1f}% (보고서: 85.7%)'.format(feb_c/feb_s*100))

ws_s = wb['2. S Plan']
gt = ws_s['X209'].value
print('S Plan Grand Total: ${:,.0f} (보고서: 50,065,917)'.format(gt))

# 월별 COGS% H1/H2 평균 검증
months = ['B','C','D','E','F','G','H','I','J','K','L','M']
labels = ['1월','2월','3월','4월','5월','6월','7월','8월','9월','10월','11월','12월']
h1_pct = []
h2_pct = []
for i, col in enumerate(months):
    s = ws[col+'4'].value
    c = ws[col+'5'].value
    pct = c/s*100
    if i < 6:
        h1_pct.append(pct)
    else:
        h2_pct.append(pct)

print('\nH1 평균 COGS%: {:.1f}% (보고서: 73.7%)'.format(sum(h1_pct)/len(h1_pct)))
print('H2 평균 COGS%: {:.1f}% (보고서: 68.6%)'.format(sum(h2_pct)/len(h2_pct)))
print('Q4 평균 COGS%: {:.1f}% (보고서: 65.7%)'.format(sum(h2_pct[3:])/3))
print('\n검증 완료.')
