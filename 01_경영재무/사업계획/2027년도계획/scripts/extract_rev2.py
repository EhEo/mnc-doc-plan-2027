# Rev2 사업계획 핵심 데이터 추출 스크립트
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from openpyxl import load_workbook
from openpyxl.chartsheet import Chartsheet

SRC = '../working/20260615-2027Y BIZ Plan_Rev2.xlsx'
wb = load_workbook(SRC, data_only=True)

print(f'시트 목록: {wb.sheetnames}\n')

# ── 1. #REF! 오류 수 확인 ────────────────────────────────────
wb_f = load_workbook(SRC, data_only=False)
ref_count = 0
ref_sheets = {}
for sn in wb_f.sheetnames:
    ws = wb_f[sn]
    if isinstance(ws, Chartsheet): continue
    cnt = sum(1 for row in ws.iter_rows() for c in row
              if c.value and '#REF!' in str(c.value))
    if cnt:
        ref_sheets[sn] = cnt
        ref_count += cnt
print(f'=== #REF! 오류 현황 ===')
print(f'  총계: {ref_count}개')
for sn, cnt in ref_sheets.items():
    print(f'  [{sn}]: {cnt}개')
print()

# ── 2. IS (손익계산서) — P열(Rev19/현계획), N열(2027합계) ────
print('===== 10. IS (손익계산서) =====')
ws_is = wb['10. IS']
for row in ws_is.iter_rows(min_row=1, max_row=35):
    vals = {}
    for c in row:
        if c.value is not None:
            vals[c.column_letter] = c.value
    if vals:
        # 주요 컬럼만 출력
        label = vals.get('A', '')
        n_val = vals.get('N', '')
        p_val = vals.get('P', '')
        q_val = vals.get('Q', '')
        if any([n_val, p_val]) and label:
            print(f'  Row{row[0].row} {label}: N={n_val} | P={p_val} | Q={q_val}')
        elif p_val and not label:
            print(f'  Row{row[0].row} (비율행): P={p_val}')
print()

# ── 3. IS 전체 헤더 확인 (어떤 컬럼이 무엇인지) ──────────────
print('===== IS Row3 헤더 =====')
ws_is3 = wb['10. IS']
for c in ws_is3[3]:
    if c.value:
        print(f'  {c.column_letter}: {repr(c.value)}')
print()

# ── 4. L PLAN — 인원 현황 ────────────────────────────────────
print('===== 4. L PLAN 인원 현황 =====')
ws_l = wb['4. L PLAN']
key_rows = [4, 13, 25, 29, 36, 40, 53, 66, 77, 79, 90, 97, 141, 142, 143]
for r in key_rows:
    row_data = {}
    for c in ws_l[r]:
        if c.value is not None:
            row_data[c.column_letter] = c.value
    if row_data:
        label = row_data.get('A', row_data.get('B', ''))
        # S(Jan.27)와 T(Feb.27) 값
        s_val = row_data.get('S', '')
        t_val = row_data.get('T', '')
        e_val = row_data.get('E', '')
        f_val = row_data.get('F', '')
        print(f'  Row{r}: A={repr(label)[:40]} | E={e_val} | F={f_val} | S={s_val} | T={t_val}')
print()

# ── 5. L PLAN a — 부서별 인원 ────────────────────────────────
print('===== 4. L PLAN a (부서별) =====')
try:
    ws_la = wb['4. L PLAN a']
    for r in range(1, 60):
        row_data = {}
        for c in ws_la[r]:
            if c.value is not None:
                row_data[c.column_letter] = c.value
        if row_data and row_data.get('A'):
            a = row_data.get('A', '')
            b = row_data.get('B', '')
            f = row_data.get('F', '')
            print(f'  Row{r}: A={repr(str(a))[:30]} | B={repr(str(b))[:20]} | F={f}')
except: print('  시트 없음')
print()

# ── 6. S Plan — 매출 수량 합계 (X열 = 2027 Total) ─────────────
print('===== 2. S Plan (X열 = 2027 Total) 주요행 =====')
ws_s = wb['2. S Plan']
for r in range(1, 200):
    row = ws_s[r]
    vals = {c.column_letter: c.value for c in row if c.value is not None}
    x = vals.get('X', '')
    a = vals.get('A', '')
    b = vals.get('B', '')
    c_val = vals.get('C', '')
    f = vals.get('F', '')
    g = vals.get('G', '')
    if x and x != vals.get('A', ''):  # X열에 값이 있는 행
        print(f'  Row{r}: A={repr(str(a))[:20]} C={repr(str(c_val))[:15]} F={repr(str(f))[:20]} G={g} X={x}')
    if r > 150:
        break
print()

# ── 7. 1. Assumption — 사업 전제 ─────────────────────────────
print('===== 1. Assumption 주요내용 =====')
ws_a = wb['1. Assumption']
for r in range(1, 40):
    row_data = {}
    for c in ws_a[r]:
        if c.value is not None:
            row_data[c.column_letter] = c.value
    if row_data:
        relevant = {k: v for k, v in row_data.items()
                   if k in ('A','B','C','D','E','F')}
        if relevant:
            print(f'  Row{r}: ' + ' | '.join(f'{k}={repr(str(v))[:40]}' for k,v in relevant.items()))
print()

# ── 8. I Plan — 투자 계획 ────────────────────────────────────
print('===== 8. I Plan 투자 계획 =====')
ws_i = wb['8. I Plan']
for r in range(1, 55):
    row_data = {}
    for c in ws_i[r]:
        if c.value is not None:
            row_data[c.column_letter] = c.value
    if row_data:
        a = row_data.get('A', '')
        b = row_data.get('B', '')
        c_v = row_data.get('C', '')
        d = row_data.get('D', '')
        e = row_data.get('E', '')
        f = row_data.get('F', '')
        g = row_data.get('G', '')
        if any([a, b, d, e]):
            print(f'  Row{r}: A={repr(str(a))[:15]} B={repr(str(b))[:15]} C={c_v} D={repr(str(d))[:35]} E={e} F={f} G={g}')
print()

# ── 9. CF — 현금흐름 주요 항목 ───────────────────────────────
print('===== CF 시트 목록 =====')
cf_sheets = [s for s in wb.sheetnames if 'CF' in s or 'cf' in s.lower()]
print(f'  CF 관련 시트: {cf_sheets}')
for sn in cf_sheets[:2]:
    ws_cf = wb[sn]
    print(f'\n  [{sn}] 주요 행:')
    for r in range(1, 55):
        row_data = {}
        for c in ws_cf[r]:
            if c.value is not None:
                row_data[c.column_letter] = c.value
        if row_data and (row_data.get('A') or row_data.get('B')):
            a = row_data.get('A', '')
            b = row_data.get('B', '')
            c_v = row_data.get('C', '')
            n_v = row_data.get('N', '') or row_data.get('O', '')
            print(f'    Row{r}: A={repr(str(a))[:25]} B={repr(str(b))[:25]} C={c_v} →N/O={n_v}')

# ── 10. Rev1 vs Rev2 IS 비교를 위한 P열 전체 ─────────────────
print('\n===== IS P열 전체값 (Rev2) =====')
ws_is2 = wb['10. IS']
for r in range(3, 35):
    a = ws_is2.cell(r, 1).value
    p = ws_is2.cell(r, 16).value
    n = ws_is2.cell(r, 14).value
    if p or n:
        print(f'  Row{r}: A={repr(str(a))[:35]} | N={n} | P={p}')
