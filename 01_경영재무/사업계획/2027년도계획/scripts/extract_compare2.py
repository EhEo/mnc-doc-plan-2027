# 2026 vs 2027 M Plan(제조원가) 세부 비교 추출
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import load_workbook
import glob as _g

SRC_26 = _g.glob('../sources/*Rev27*.xlsx') + _g.glob('../sources/*Rev27*.XLSX')
SRC_26 = SRC_26[0]
SRC_27 = '../working/20260615-2027Y BIZ Plan_Rev2.xlsx'

wb26 = load_workbook(SRC_26, data_only=True)
wb27 = load_workbook(SRC_27, data_only=True)

def get_col_sum(ws, row, start_col='B', end_col='M'):
    """월별(B~M) 합계 계산 — N열 캐시값이 없을 때 수동 합산"""
    total = 0
    for col in ws[row]:
        if start_col <= col.column_letter <= end_col:
            if isinstance(col.value, (int, float)):
                total += col.value
    return total

# ── 2026 M Plan 전체 구조 파악 ───────────────────────────
print('===== [2026 Rev27] 5. M Plan 전체 구조 =====')
ws_m26 = wb26['5. M Plan']
# 헤더 행 확인
for r in range(1, 10):
    row_data = {c.column_letter: c.value for c in ws_m26[r] if c.value is not None}
    if row_data:
        print(f'  Row{r}: ' + ' | '.join(f'{k}={repr(str(v))[:30]}' for k,v in list(row_data.items())[:8]))

print()
# 전체 A열 항목 나열 (N열 = 연간 합계)
for r in range(1, 150):
    row_data = {c.column_letter: c.value for c in ws_m26[r] if c.value is not None}
    a = row_data.get('A', '')
    b = row_data.get('B', '')
    n = row_data.get('N', '')
    if a and str(a).strip():
        print(f'  Row{r:3d}: A={repr(str(a))[:45]} | N={n}')

# ── 2027 M Plan 전체 구조 ────────────────────────────────
print('\n\n===== [2027 Rev2] 5. M Plan 전체 구조 =====')
ws_m27 = wb27['5. M Plan']
for r in range(1, 150):
    row_data = {c.column_letter: c.value for c in ws_m27[r] if c.value is not None}
    a = row_data.get('A', '')
    n = row_data.get('N', '')
    if a and str(a).strip():
        print(f'  Row{r:3d}: A={repr(str(a))[:45]} | N={n}')

# ── IS Row6 Write-off 월별 확인 ──────────────────────────
print('\n\n===== [2026 Rev27] IS Row6 Write-off 월별 =====')
ws_is26 = wb26['10. IS']
row6 = {c.column_letter: round(c.value,0) if isinstance(c.value,float) else c.value
        for c in ws_is26[6] if c.value is not None}
for k,v in row6.items():
    print(f'  {k}: {v}')

# ── 2026 AP Plan (간접비/경비) ───────────────────────────
print('\n\n===== [2026 Rev27] AP Plan (간접비) 주요행 =====')
if 'AP Plan' in wb26.sheetnames:
    ws_ap26 = wb26['AP Plan']
    for r in range(1, 80):
        row_data = {c.column_letter: c.value for c in ws_ap26[r] if c.value is not None}
        a = row_data.get('A','')
        n = row_data.get('N','')
        if a and str(a).strip():
            print(f'  Row{r:3d}: A={repr(str(a))[:45]} | N(2026합계)={n}')

# ── 2027 AP Plan ─────────────────────────────────────────
print('\n\n===== [2027 Rev2] AP Plan (간접비) 주요행 =====')
if 'AP Plan' in wb27.sheetnames:
    ws_ap27 = wb27['AP Plan']
    for r in range(1, 80):
        row_data = {c.column_letter: c.value for c in ws_ap27[r] if c.value is not None}
        a = row_data.get('A','')
        n = row_data.get('N','')
        if a and str(a).strip():
            print(f'  Row{r:3d}: A={repr(str(a))[:45]} | N(2027합계)={n}')

# ── 2026 L Plan a — 인원절감 섹션 ────────────────────────
print('\n\n===== [2026 Rev27] L PLAN a 인원절감 섹션 =====')
ws_la26 = wb26['4. L PLAN a']
for r in range(38, 60):
    row_data = {c.column_letter: c.value for c in ws_la26[r] if c.value is not None}
    if row_data:
        line = ' | '.join(f'{k}={repr(str(v))[:25]}' for k,v in list(row_data.items())[:8])
        print(f'  Row{r}: {line}')

# ── 2026 S Plan 고객사별 상세 (단가·수량 포함) ──────────
print('\n\n===== [2026 Rev27] S Plan 고객사별 상세 매출 =====')
ws_s26 = wb26['2. S Plan']
for r in range(108, 210):
    row_data = {c.column_letter: c.value for c in ws_s26[r] if c.value is not None}
    if not row_data:
        continue
    a = row_data.get('A','')
    c_v = row_data.get('C','')
    f = row_data.get('F','')
    g = row_data.get('G','')
    x = row_data.get('X','')
    if x or 'Total' in str(c_v) or 'Total' in str(f) or a:
        print(f'  Row{r}: A={repr(str(a))[:18]} C={repr(str(c_v))[:18]} F={repr(str(f))[:18]} G={g} X={x}')

print('\n===== 완료 =====')
