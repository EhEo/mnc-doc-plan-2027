# Excel KPI 수립서 → JSON 변환 스크립트

import json, argparse
from pathlib import Path
from datetime import datetime
import openpyxl

class ExcelParseError(Exception):
    pass

DEPT_MAP = {
    "관리부문":"관리부문", "생산1부문":"생산1부문",
    "생산2부문":"생산2부문", "품질부문":"품질부문",
    "개발영업부문":"개발영업부문"
}

def parse_kpi_excel(xlsx_path: Path, year: str = "2026", half: str = "H2") -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "KPI목록" not in wb.sheetnames:
        raise ExcelParseError("'KPI목록' 시트가 없습니다")
    ws = wb["KPI목록"]

    dept_raw = ws["B2"].value or ""
    dept = DEPT_MAP.get(str(dept_raw).strip())
    if not dept:
        raise ExcelParseError(f"부문명 인식 불가: '{dept_raw}'")

    def _num(val: object) -> float:
        try:
            return float(val) if val is not None else 0.0  # type: ignore[arg-type]
        except (TypeError, ValueError):
            return 0.0

    kpis = []
    for row in ws.iter_rows(min_row=4, max_row=13, values_only=True):
        row_list = (list(row) + [None]*10)[:10]
        kpi_id, category, name, target, unit, weight, method, owner, deadline, _ = row_list
        if not kpi_id or not name:
            continue
        kpis.append({
            "id": str(kpi_id).strip(),
            "category": str(category or "").strip(),
            "name": str(name).strip(),
            "target": _num(target),
            "unit": str(unit or "").strip(),
            "weight": _num(weight),
            "measurement_method": str(method or "").strip(),
            "owner": str(owner or "").strip(),
            "deadline": str(deadline or "").strip(),
            "monthly_actuals": {}
        })

    return {
        "dept": dept, "year": year, "half": half,
        "version": "v1.0", "status": "draft",
        "generated_at": datetime.now().strftime("%Y-%m-%d"),
        "kpis": kpis
    }

def main():
    parser = argparse.ArgumentParser(description="Excel KPI 수립서 → JSON 변환")
    parser.add_argument("xlsx", help="입력 Excel 경로")
    parser.add_argument("--year", default="2026")
    parser.add_argument("--half", default="H2")
    parser.add_argument("--out", help="출력 JSON 경로 (기본: 입력 파일과 같은 폴더)")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    data = parse_kpi_excel(xlsx_path, args.year, args.half)

    out_path = Path(args.out) if args.out else xlsx_path.parent / "kpi_data.json"
    out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"✅ 변환 완료: {out_path}")

if __name__ == "__main__":
    main()
