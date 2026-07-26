from pathlib import Path

from u9c_catalog.doc_generator import generate_html, generate_excel
from u9c_catalog.extractor import ExtractResult
from u9c_catalog.noise_classifier import classify_all
from tests.fixtures.sample_metadata import make_receivement_table, make_trl_table


def _result():
    r = ExtractResult(tables=[make_receivement_table(), make_trl_table()])
    classify_all(r.tables)
    return r


def test_generate_html(tmp_path):
    out = tmp_path / "dd.html"
    generate_html(_result(), str(out))
    html = out.read_text(encoding="utf-8")
    assert "PM_Receivement" in html
    assert "DocNo" in html
    # 시스템 컬럼은 표시하되 구분 표기가 있어야 함
    assert "시스템" in html
    # 부수 테이블 표기
    assert "부수" in html


def test_generate_excel(tmp_path):
    out = tmp_path / "dd.xlsx"
    generate_excel(_result(), str(out))
    assert Path(out).exists()
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert "Objects" in wb.sheetnames
    assert "Columns" in wb.sheetnames
