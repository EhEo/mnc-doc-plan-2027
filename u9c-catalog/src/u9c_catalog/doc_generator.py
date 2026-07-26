# 메타데이터로 HTML/Excel 데이터 사전을 생성 (한글 폰트 맑은 고딕)
import json
from jinja2 import Environment, PackageLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Font

from u9c_catalog.extractor import ExtractResult

_env = Environment(
    loader=PackageLoader("u9c_catalog", "templates"),
    autoescape=select_autoescape(["html"]),
)


def generate_html(result: ExtractResult, out_path: str) -> None:
    tmpl = _env.get_template("data_dictionary.html.j2")
    html = tmpl.render(tables=result.tables)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)


def generate_excel(result: ExtractResult, out_path: str) -> None:
    wb = Workbook()
    ws_obj = wb.active
    ws_obj.title = "Objects"
    ws_obj.append(["schema", "object", "type", "row_count", "is_auxiliary", "aux_reason"])
    for t in result.tables:
        ws_obj.append([t.schema, t.name, t.object_type, t.row_count,
                       "Y" if t.is_auxiliary else "N", t.auxiliary_reason or ""])

    ws_col = wb.create_sheet("Columns")
    ws_col.append(["schema", "object", "column", "type", "nullable", "pk", "is_system", "system_reason"])
    for t in result.tables:
        for c in t.columns:
            ws_col.append([t.schema, t.name, c.name, c.type_display or c.data_type,
                           "Y" if c.is_nullable else "N", "PK" if c.is_pk else "",
                           "Y" if c.is_system else "N", c.system_reason or ""])

    for ws in (ws_obj, ws_col):
        for cell in ws[1]:
            cell.font = Font(name="맑은 고딕", bold=True)

    wb.save(out_path)


def filter_priority_tables(tables, priorities):
    """is_priority=True인 테이블만 원래 순서로 반환."""
    pri = {(p.schema, p.name) for p in priorities if p.is_priority}
    return [t for t in tables if (t.schema, t.name) in pri]


def generate_priority_json(priorities, out_path):
    """우선순위 랭킹을 다운스트림(MES/대시보드)용 JSON으로 출력 (rank 오름차순)."""
    data = [
        {"object": f"{p.schema}.{p.name}", "rank": p.rank, "score": round(p.score, 3),
         "is_priority": p.is_priority, "reason": p.reason}
        for p in sorted(priorities, key=lambda x: x.rank)
    ]
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def generate_domain_map_json(assignments, edges, out_path):
    """도메인별 테이블 + 관계 엣지를 MES/대시보드용 JSON으로 내보낸다."""
    domains: dict = {}
    for a in assignments:
        if a.domain == "미분류":
            continue
        d = domains.setdefault(a.domain, {"tables": []})
        d["tables"].append({"object": a.full_name, "role": a.role,
                            "confidence": round(a.confidence, 2), "evidence": a.evidence,
                            "verified_by": a.verified_by})
    rels = [{"from": e.from_object, "column": e.from_column, "to": e.to_object,
             "kind": e.kind, "evidence": e.evidence} for e in edges]
    data = {"domains": domains, "relations": rels}
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
