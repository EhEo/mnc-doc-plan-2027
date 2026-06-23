# AS9100D 신입사원 품질교육 워드 문서(4파트)를 생성하는 스크립트
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_korean_font(element, font_name="맑은 고딕"):
    rpr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        element.insert(0, rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ["ascii", "hAnsi", "eastAsia", "cs"]:
        rfonts.set(qn(f"w:{attr}"), font_name)


def apply_korean_fonts(doc):
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        try:
            st = doc.styles[style_name]
            st.font.name = "맑은 고딕"
            set_korean_font(st.element.find(qn("w:rPr")) or st.element, "맑은 고딕")
        except Exception:
            pass


def add_heading(doc, text, level, color=None):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if color:
            run.font.color.rgb = color
    return p


def add_body(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.8)
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10.5)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.font.name = "맑은 고딕"
        run.font.bold = True
        run.font.size = Pt(10)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, cell_text in enumerate(row_data):
            row_cells[c_idx].text = cell_text
            run = row_cells[c_idx].paragraphs[0].runs[0]
            run.font.name = "맑은 고딕"
            run.font.size = Pt(10)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.add_paragraph()


def add_box_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.right_indent = Cm(1)
    run = p.add_run(f"📌  {text}")
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10)
    run.font.italic = True
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return p


# ── 문서 생성 시작 ──────────────────────────────────────────────────
doc = Document()
apply_korean_fonts(doc)

# 페이지 여백
sec = doc.sections[0]
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3)
sec.right_margin = Cm(2.5)

BLUE  = RGBColor(0x1F, 0x49, 0x7D)
GREEN = RGBColor(0x37, 0x86, 0x46)

# ════════════════════════════════════════════════════════════════════
# 표지
# ════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("(주)에이로텍")
r.font.name = "맑은 고딕"; r.font.size = Pt(16); r.font.bold = True
r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

doc.add_paragraph()

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = t2.add_run("신입사원 품질교육 교재")
r2.font.name = "맑은 고딕"; r2.font.size = Pt(28); r2.font.bold = True
r2.font.color.rgb = BLUE
r2._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

t3 = doc.add_paragraph()
t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = t3.add_run("AS9100D Rev D 기반 항공우주 품질경영시스템")
r3.font.name = "맑은 고딕"; r3.font.size = Pt(14)
r3._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

doc.add_paragraph()
doc.add_paragraph()

meta_data = [
    ("작성일", "2026년 5월 22일"),
    ("작성 부서", "품질경영팀"),
    ("적용 대상", "전 신입사원"),
    ("적용 기준", "AS9100D Rev D / AS9102 Rev B / AS9146 / 위조품방지절차(QP-841)"),
]
for label, value in meta_data:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l = p.add_run(f"{label}: "); r_l.font.name = "맑은 고딕"; r_l.font.size = Pt(11); r_l.bold = True
    r_l._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r_v = p.add_run(value); r_v.font.name = "맑은 고딕"; r_v.font.size = Pt(11)
    r_v._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# PART 1 — 품질의 개념과 AS9100D 개요
# ════════════════════════════════════════════════════════════════════
add_heading(doc, "PART 1. 품질의 개념과 AS9100D 개요", 1, BLUE)

add_heading(doc, "1.1 항공우주 품질이란 무엇인가", 2)
add_body(doc, "품질이란 '요구사항을 충족하는 정도'를 의미합니다. 항공우주 산업에서 품질은 단순한 제품 수준을 넘어 인명 안전과 직결됩니다.")
add_bullet(doc, "작은 결함 하나가 대형 항공 사고로 이어질 수 있는 고위험 산업")
add_bullet(doc, "Zero Defect(무결점) 원칙: 허용 가능한 불량률이란 없다")
add_bullet(doc, "제품 수명 주기 전반에 걸친 품질 보증 요구")
add_box_note(doc, "실제 사례: 1988년 알로하 항공 사고는 정비 기록 오류와 금속 피로 검사 불량이 복합적으로 작용한 결과였습니다.")

doc.add_paragraph()
add_heading(doc, "1.2 AS9100D 규격의 탄생 배경과 구조", 2)
add_body(doc, "AS9100D는 ISO 9001:2015를 기반으로 항공·우주·방산 분야의 추가 요구사항을 결합한 국제 표준 규격입니다.")
add_table(doc,
    ["구분", "내용"],
    [
        ["제정 기관", "SAE International / IAQG (국제항공우주품질그룹)"],
        ["기반 규격", "ISO 9001:2015"],
        ["최신 버전", "AS9100 Rev D (2016 발행)"],
        ["조항 구성", "4~10조 (총 7개 조항 + 항공우주 특수 요구사항)"],
        ["주요 대상", "항공기 제조, 부품 공급업체, MRO, 방산업체"],
    ]
)

add_heading(doc, "1.3 프로세스 접근법과 리스크 기반 사고", 2)
add_body(doc, "AS9100D는 모든 업무를 '프로세스'의 관점에서 바라볼 것을 요구합니다.")
add_bullet(doc, "프로세스 접근법: 입력 → 활동 → 출력의 흐름으로 업무를 설계")
add_bullet(doc, "리스크 기반 사고: 문제가 생기기 전에 '무엇이 잘못될 수 있는가'를 먼저 분석")
add_bullet(doc, "FMEA(고장모드영향분석): 잠재적 결함과 그 영향을 사전에 평가하는 도구")
add_bullet(doc, "기회(Opportunity): 리스크 관리를 통해 개선의 기회도 함께 식별")
doc.add_paragraph()
add_body(doc, "항공우주 리스크 관리의 4단계 연결 흐름 (신입사원 필수 이해 사항).")
add_table(doc,
    ["단계", "개념", "의미"],
    [
        ["1단계", "특별요구사항 (Special Requirements, SR)", "고객 또는 규제기관이 지정한 특수 기술·검사 요구사항 (예: 표면처리 사양, 특정 시험 횟수)"],
        ["2단계", "중요품목 (Critical Items, CI)", "결함 발생 시 안전 또는 임무에 직접 영향을 미치는 부품·공정 (비행 안전 부품 포함)"],
        ["3단계", "핵심특성 (Key Characteristics, KC)", "CI에서 도출된 치수·재료·공정 파라미터 중 품질에 결정적인 변수 (예: 특정 볼트 조임 토크값)"],
        ["4단계", "관리계획·검사강도", "KC별로 측정 방법·주기·합부판정 기준을 관리계획서에 명시하고 강화된 검사 적용"],
    ]
)
add_box_note(doc, "현장 실무 팁: 작업지시서에 '★ KC' 또는 '★ CI' 마크가 있는 항목은 일반 공정보다 훨씬 엄격한 측정과 기록이 요구됩니다. 절대 건너뛰지 마세요.")

add_heading(doc, "1.4 고객 요구사항과 법적 규제 준수", 2)
add_body(doc, "항공우주 공급업체는 고객사 요구사항과 각국 항공 당국의 규정을 동시에 준수해야 합니다.")
add_bullet(doc, "주요 고객사 요구: Boeing D6-82479, Airbus AIPI 등 고객별 추가 요구사항 존재")
add_bullet(doc, "국내외 규제 기관: 국토교통부(MOLIT), FAA(미국), EASA(유럽) 인증 요구")
add_bullet(doc, "수출 통제: ITAR(국제무기거래규정) 및 EAR 준수 의무")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# PART 2 — 품질경영시스템 기본
# ════════════════════════════════════════════════════════════════════
add_heading(doc, "PART 2. 품질경영시스템 기본", 1, BLUE)
add_body(doc, "품질경영시스템(QMS)은 일관된 품질을 보장하기 위해 조직이 운영하는 체계입니다. 이 파트에서는 문서 관리와 절차 준수의 기본을 학습합니다.")

add_heading(doc, "2.1 문서 체계 이해", 2)
add_body(doc, "우리 회사의 문서는 4계층으로 구성됩니다.")
add_table(doc,
    ["계층", "문서 유형", "문서 번호", "역할"],
    [
        ["1계층", "품질매뉴얼", "QM-001", "AS9100D 전 조항 포괄, 최상위 방침 선언"],
        ["2계층", "절차서", "QP-xxx", "업무 흐름 및 책임자 정의 (30건)"],
        ["3계층", "지침서", "WI-xxx", "현장 작업 세부 방법 (15건)"],
        ["4계층", "양식·기록", "QF-xxx", "증거로서 보관되는 기록 (21건)"],
    ]
)
add_box_note(doc, "반드시 기억하세요: 작업 전 반드시 최신 개정본(Revision)을 확인해야 합니다. 구버전 문서로 작업하면 부적합 판정을 받습니다.")

add_heading(doc, "2.2 문서화된 정보의 관리", 2)
add_bullet(doc, "배포된 문서의 개정 상태(Rev.) 및 발행일 확인 필수")
add_bullet(doc, "승인되지 않은 비공식 문서(메모지, 개인 노트) 사용 금지")
add_bullet(doc, "폐지된 구버전 문서는 즉시 회수 또는 '폐기' 스탬프 날인")
add_bullet(doc, "전자 문서의 경우 DMS(문서관리시스템) 상의 최신본만 유효")

add_heading(doc, "2.3 표준작업지침서(SOP) 준수", 2)
add_body(doc, "'결정된 대로 작업하기'는 품질의 기본 원칙입니다.")
add_bullet(doc, "SOP에 정의된 순서, 공구, 재료를 그대로 따를 것")
add_bullet(doc, "임의 판단에 의한 공정 변경 절대 금지 — 반드시 엔지니어링 변경 절차(ECN) 거칠 것")
add_bullet(doc, "작업 중 문제 발생 시: 작업 중지 → 상급자 보고 → 지시에 따라 조치")
add_box_note(doc, "SOP 이탈은 단순 실수가 아닌 '시스템 위반'으로 분류되며 CAR(시정조치 요청) 발행 사유가 됩니다.")

add_heading(doc, "2.4 기록의 무결성(Data Integrity)", 2)
add_body(doc, "항공우주 산업에서 기록은 법적 효력을 가지는 증거입니다.")
add_table(doc,
    ["올바른 기록 방법", "금지 사항"],
    [
        ["사실에 기반하여 실시간 기록", "수정액(화이트) 사용"],
        ["수정 시 한 줄 긋고 서명·날짜", "기억에 의존한 사후 기록"],
        ["볼펜으로 선명하게 작성", "연필 또는 지워지는 펜 사용"],
        ["규정된 보존 기한 준수", "임의 폐기"],
    ]
)

add_heading(doc, "2.5 내부심사와 시정조치(CAR)", 2)
add_bullet(doc, "내부심사: 조직 스스로 규정 준수 여부를 점검하는 활동 (연 1회 이상)")
add_bullet(doc, "부적합(NC) 발견 시 CAR 발행 → 근본원인 분석(5Why/Fish-bone) → 시정조치 실행")
add_bullet(doc, "CAR의 세 가지 유형: Major(중결함), Minor(경결함), OFI(개선 기회)")
add_bullet(doc, "신입사원도 부적합 사항을 발견하면 즉시 보고할 의무가 있습니다")
doc.add_paragraph()
add_body(doc, "AS9100D가 강조하는 Human Factors 기반 근본원인 분석 — 기술적 원인 외에 아래 인적 요인을 반드시 검토합니다.")
add_table(doc,
    ["Human Factors 항목", "체크 질문"],
    [
        ["피로 / 수면 부족", "장시간 연속 작업이나 교대 직후 발생했는가?"],
        ["주의분산 / 중단", "작업 도중 외부 요인으로 집중이 끊겼는가?"],
        ["교육·훈련 부족", "해당 작업에 대한 교육이 충분히 이루어졌는가?"],
        ["작업환경 문제", "조명, 소음, 공간이 작업자에게 불리했는가?"],
        ["의사소통 오류", "구두·서면 지시가 불명확하거나 언어 장벽이 있었는가?"],
    ]
)
add_box_note(doc, "재발방지 대책은 '주의를 기울이겠습니다'처럼 개인 의지에 의존해서는 안 됩니다. 반드시 공정·절차·환경 개선으로 연결되어야 합니다.")

add_heading(doc, "2.6 구성관리(Configuration Management, CM)", 2)
add_body(doc, "구성관리는 제품·공정·문서가 서로 정확하게 일치하도록 유지하는 체계입니다. (AS9100D 8.1.2)")
add_bullet(doc, "구성 기준선(Configuration Baseline): 특정 시점에 승인된 도면 Rev, BOM, 작업지시서, NC 프로그램, 자재 LOT의 조합")
add_bullet(doc, "작업 전 일치 확인: 도면 Rev와 작업지시서 Rev가 동일한지 반드시 대조")
add_bullet(doc, "변경 통제(Change Control): 도면·BOM·공정 중 하나라도 변경되면 CCB(형상통제위원회) 검토 후 승인 필요")
add_bullet(doc, "현장 금지 사항: 승인되지 않은 도면 개정, 구두 지시만으로 BOM 변경, 구버전 NC 프로그램 사용")
add_table(doc,
    ["구성요소", "확인 내용", "불일치 시 조치"],
    [
        ["도면 (Drawing)", "Rev 레벨·발행일 일치", "작업 중지 → 품질팀 확인"],
        ["BOM (자재명세서)", "부품번호·수량·소재 일치", "자재 교체 후 재확인"],
        ["작업지시서 (WI/SOP)", "Rev 및 공정 순서 일치", "최신본으로 교체 후 진행"],
        ["NC 프로그램", "버전 번호 및 검증 이력 확인", "승인된 버전으로 교체"],
    ]
)
add_box_note(doc, "실제 부적합 사례: 도면은 Rev D인데 작업지시서가 Rev C인 상태로 100개를 가공한 후 전량 재검사를 실시한 사례가 있습니다. 작업 시작 전 30초 확인이 몇 시간의 손실을 막습니다.")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# PART 3 — 현장 품질 실무
# ════════════════════════════════════════════════════════════════════
add_heading(doc, "PART 3. 현장 품질 실무", 1, BLUE)
add_body(doc, "현장에서 직접 마주하게 될 품질 실무 절차들을 학습합니다. 이 항목들은 AS9100D 심사에서 가장 자주 지적되는 고위험 영역입니다.")

add_heading(doc, "3.1 FAI — 초도품 검사 (AS9102 기준)", 2)
add_body(doc, "FAI(First Article Inspection)는 신규 부품 또는 공정 변경 후 최초 생산품의 모든 요구사항을 100% 검증하는 절차입니다.")
add_bullet(doc, "적용 시점: 신규 부품 첫 생산, 설계 변경, 공정/장비/재료 변경, 생산 중단 후 재개(2년 이상)")
add_bullet(doc, "FAI 구성 요소: ① 설계 문서 검토, ② 자재 성적서 확인, ③ 치수 검사(전수), ④ 기능 시험, ⑤ FAIR 보고서 작성")
add_bullet(doc, "AS9102 Rev B(또는 적용 개정판) 기준으로 FAIR(First Article Inspection Report)를 작성하며, 고객 제출·승인 여부는 계약 요구사항에 따라 결정됨")
add_box_note(doc, "FAI는 '한 번 하면 끝'이 아닙니다. 설계 또는 공정의 중요 변경이 발생할 때마다 부분 또는 전체 재실시가 요구됩니다.")

add_heading(doc, "3.2 FOD — 이물질 방지 프로그램 (AS9146 기준)", 2)
add_body(doc, "FOD(Foreign Object Damage/Debris)는 항공기 내부에 유입된 이물질로 인한 손상 또는 그 이물질 자체를 의미합니다.")
add_bullet(doc, "FOD 통제 구역(FOD Zone): 지정된 청정 작업 구역, 출입 시 공구·부품 수량 기록 필수")
add_bullet(doc, "공구 관리(Tool Control): 사용 전·후 공구 수량 카운트, 분실 시 즉각 보고")
add_bullet(doc, "마무리 전 점검: 조립 완료 전 내부 이물질 잔류 여부 육안 검사")
add_bullet(doc, "위반 사례: 장갑, 볼트, 드릴 비트 등이 기체 내부에서 발견되면 전수 재검사 대상")
add_box_note(doc, "FOD 1건이 항공기 전체의 비행 정지로 이어질 수 있습니다. 작업 후 '빠진 것이 없는가'를 항상 확인하세요.")

add_heading(doc, "3.3 식별 및 추적성(Traceability) 관리", 2)
add_body(doc, "항공우주 부품은 원재료부터 완제품까지 모든 이력이 추적 가능해야 합니다.")
add_bullet(doc, "LOT/배치 번호: 동일 조건에서 생산된 부품 묶음의 고유 식별 번호")
add_bullet(doc, "시리얼 번호(S/N): 개별 부품의 고유 번호 (고안전 부품에 필수)")
add_bullet(doc, "자재 성적서(Material Certificate): 원재료 화학 성분, 기계적 특성 증명 문서")
add_bullet(doc, "부적합 발생 시 해당 LOT 전체 격리 및 역추적으로 영향 범위 확인")
add_table(doc,
    ["추적성 요소", "기록 내용", "보존 기간"],
    [
        ["자재 입고 기록", "공급업체, LOT번호, 성적서 번호", "10년 이상"],
        ["공정 기록", "작업자, 장비번호, 날짜, 조건", "10년 이상"],
        ["검사 기록", "검사원, 계측기 번호, 판정 결과", "10년 이상"],
        ["출하 기록", "출하 일자, 수량, 고객 PO 번호", "10년 이상"],
    ]
)

add_heading(doc, "3.4 부적합품(NC) 처리 절차", 2)
add_body(doc, "부적합품을 발견하면 아래 순서에 따라 처리합니다.")
steps = [
    "① 발견 즉시 작업 중지",
    "② 빨간 '부적합품(NC)' 태그 부착",
    "③ 지정된 격리 구역(Red Zone)으로 이동",
    "④ 부적합 보고서 작성 및 품질팀 통보",
    "⑤ 처분 결정: 수리 / 재작업 / 특채 / 폐기",
    "⑥ 처분 완료 후 기록 마감",
]
for step in steps:
    add_bullet(doc, step)
add_box_note(doc, "절대 금지: 부적합품을 정상품과 섞어 놓거나, 태그 없이 방치하거나, 임의로 폐기하는 행위")

add_heading(doc, "3.5 계측기 관리 및 교정(Calibration)", 2)
add_bullet(doc, "사용 전 교정 유효기간 확인 (교정 스티커의 유효일 확인)")
add_bullet(doc, "교정 기한 초과 계측기 사용 절대 금지 — 즉시 교정 의뢰")
add_bullet(doc, "낙하·충격 발생 계측기는 즉시 재교정 의뢰 후 사용")
add_bullet(doc, "교정된 계측기라도 사용 전 영점(Zero) 확인 습관화")

add_heading(doc, "3.6 위조품 및 의심 위조품 방지 (AS9100D 8.1.4)", 2)
add_body(doc, "위조품(Counterfeit Part)은 원산지·성능·사양을 속인 불법 복제 부품입니다. 항공우주 산업에서 위조품 유입은 인명 사고로 직결됩니다.")
add_bullet(doc, "위조품 유형: 외관 위조(마킹 변조), 재생품을 신품으로 포장, 규격 미달 부품에 인증 마크 부착")
add_bullet(doc, "구매처 검증: 승인 공급업체 목록(ASL)에 등재된 업체에서만 자재 구매 — 비승인처 구매 절대 금지")
add_bullet(doc, "입고 검사 시 확인 사항: 라벨·마킹·로고 인쇄 품질, 시리얼 번호 일련성, 포장 상태, 자재 성적서 원본 여부")
doc.add_paragraph()
add_body(doc, "의심 위조품(Suspect Counterfeit) 발견 시 처리 절차.")
add_table(doc,
    ["단계", "조치 내용"],
    [
        ["① 즉시 격리", "의심 부품에 '의심 위조품(SUS-CNT)' 태그 부착, 별도 격리 구역 보관"],
        ["② 사용 금지", "해당 LOT 전체 사용 중단 — 이미 사용된 경우 구현 제품 역추적"],
        ["③ 품질팀 즉시 보고", "구매 출처·수량·식별정보를 포함한 의심 보고서(SCR) 작성"],
        ["④ 공급업체 통보", "해당 공급업체에 증거 자료와 함께 서면 통보"],
        ["⑤ 당국 보고(필요 시)", "GIDEP(미국) 또는 국가항공안전기관 보고 — 계약 요구사항 확인"],
    ]
)
add_box_note(doc, "의심스러우면 쓰지 마세요. '아마 괜찮겠지'라는 판단 한 번이 회사 전체의 인증 취소로 이어질 수 있습니다.")

doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# PART 4 — 품질 문화와 신입사원의 역할
# ════════════════════════════════════════════════════════════════════
add_heading(doc, "PART 4. 품질 문화와 신입사원의 역할", 1, BLUE)
add_body(doc, "품질은 품질팀만의 일이 아닙니다. 조직의 모든 구성원이 품질의 주인공입니다.")

add_heading(doc, "4.1 제품 안전 인식(Product Safety)", 2)
add_body(doc, "내가 만든 부품은 수백 명이 탑승한 항공기에 사용됩니다.", bold=True)
add_bullet(doc, "비행 안전 부품(Flight Safety Part): 파손 시 항공기 추락으로 이어질 수 있는 최고 위험 등급 부품")
add_bullet(doc, "부품 안전 등급(Safety Classification)에 따라 검사 강도와 기록 요건이 달라짐")
add_bullet(doc, "내 손을 거친 부품이 여객기, 전투기, 위성에 장착된다는 책임감을 항상 유지")
add_box_note(doc, "품질에 의심이 들면 즉시 멈추세요. '괜찮겠지'라는 판단은 비행 안전에 치명적입니다.")

add_heading(doc, "4.2 윤리적 행동과 성실성(Integrity)", 2)
add_body(doc, "항공우주 산업 종사자에게 요구되는 가장 중요한 자질은 정직입니다.")
add_table(doc,
    ["윤리적 행동", "금지 행동"],
    [
        ["실수 즉시 보고", "실수 은폐 또는 허위 보고"],
        ["측정값 있는 그대로 기록", "합격 범위에 맞춰 수치 조작"],
        ["불확실하면 물어보기", "확인 없이 추측으로 진행"],
        ["문제 발견 시 즉각 상급자 보고", "책임 회피 목적의 묵인"],
    ]
)
add_box_note(doc, "기록 조작은 단순 징계를 넘어 형사 처벌 대상이 될 수 있습니다. 항공 규제 당국은 기록 무결성 위반을 매우 엄중히 다룹니다.")

add_heading(doc, "4.3 지속적 개선과 PDCA 사이클", 2)
add_body(doc, "문제를 발견하고 개선하는 사람이 진정한 품질 전문가입니다.")
add_bullet(doc, "Plan(계획): 개선 목표 설정 및 방법 계획")
add_bullet(doc, "Do(실행): 소규모로 시험 적용")
add_bullet(doc, "Check(확인): 결과 측정 및 목표 달성 여부 평가")
add_bullet(doc, "Act(조치): 성공 시 표준화, 실패 시 재계획")
add_bullet(doc, "개선 제안 제도(Kaizen): 현장에서 발견한 낭비·불편 사항을 제안서로 제출")

add_heading(doc, "4.4 의사소통과 보고 체계", 2)
add_body(doc, "이상 징후는 초기에 보고할수록 조치 비용이 줄어듭니다.")
add_bullet(doc, "3단계 보고 원칙: 발견 → 1시간 내 구두 보고 → 24시간 내 서면 보고")
add_bullet(doc, "보고 내용: 무엇을 / 언제 / 어디서 / 어떻게 발견했는가")
add_bullet(doc, "Stop Work Authority(작업 중지 권한): 누구든 안전 위협 상황에서는 작업을 즉시 중지할 권한 보유")
add_bullet(doc, "열린 문화(Open Culture): 문제 보고를 장려하고 보고자를 보호하는 조직 문화")

add_heading(doc, "4.5 신입사원 첫 3개월 품질 실천 체크리스트", 2)
add_table(doc,
    ["항목", "실천 내용", "확인"],
    [
        ["문서 숙지", "담당 업무 관련 절차서(QP) 및 지침서(WI) 필독", "□"],
        ["FOD 훈련", "FOD 예방 교육 이수 및 서명", "□"],
        ["계측기 교육", "사용 계측기 교정 상태 확인 방법 숙지", "□"],
        ["기록 작성", "검사 기록지 작성 연습 및 피드백 수령", "□"],
        ["보고 연습", "부적합 발생 시뮬레이션 훈련 참가", "□"],
        ["제안 제출", "개선 제안 1건 이상 제출", "□"],
    ]
)

doc.add_paragraph()
add_heading(doc, "마무리 — 品質은 습관입니다", 2, GREEN)
add_body(doc,
    "AS9100D는 규격집이 아닙니다. 매일의 작은 실천이 쌓여 품질 문화가 됩니다. "
    "여러분 한 명 한 명이 (주)에이로텍의 품질을 만들어가는 주인공입니다.",
    bold=False
)
doc.add_paragraph()
p_final = doc.add_paragraph()
p_final.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f = p_final.add_run("\"품질은 검사를 통해 만들어지는 것이 아니라, 공정에 내재되어 있는 것이다.\" — W. Edwards Deming")
r_f.font.name = "맑은 고딕"; r_f.font.size = Pt(10); r_f.italic = True
r_f._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

# ── 저장 ─────────────────────────────────────────────────────────
output_path = r"C:\Users\MISTOP\Documents\02_AI디지털전환\AI-Data\김진욱\신입사원_품질교육_AS9100D.docx"
doc.save(output_path)
print(f"저장 완료: {output_path}")
